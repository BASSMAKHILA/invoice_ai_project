from fastapi import FastAPI, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, FileResponse
from pymongo import MongoClient
from bson import ObjectId
from datetime import datetime
import os
import shutil
import pytesseract
from PIL import Image
from pdf2image import convert_from_path
import openpyxl
from docx import Document
from PyPDF2 import PdfReader

# ----------------- Config -----------------
app = FastAPI()
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],  # limiter en production
    allow_methods=["*"],
    allow_headers=["*"]
)

client = MongoClient("mongodb://localhost:27017/")
db = client["gestion_documents"]
factures_col = db["factures"]
bdcs_col = db["bdcs"]

UPLOAD_FOLDER = "./uploads"
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

# ----------------- Helpers -----------------
def save_file(file: UploadFile):
    ext = os.path.splitext(file.filename)[1]
    file_path = os.path.join(UPLOAD_FOLDER, f"{datetime.utcnow().timestamp()}_{file.filename}")
    with open(file_path, "wb") as f:
        shutil.copyfileobj(file.file, f)
    return file_path

def extract_text(file_path: str):
    ext = os.path.splitext(file_path)[1].lower()
    text = ""
    try:
        if ext in [".png", ".jpg", ".jpeg", ".tiff", ".bmp", ".gif"]:
            text = pytesseract.image_to_string(Image.open(file_path), lang="fra")
        elif ext == ".pdf":
            # Essayer d'extraire le texte directement
            try:
                reader = PdfReader(file_path)
                for page in reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
            except Exception:
                pass
            # Si texte vide, utiliser OCR
            if not text.strip():
                pages = convert_from_path(file_path)
                for page in pages:
                    text += pytesseract.image_to_string(page, lang="fra") + "\n"
        elif ext in [".txt", ".csv"]:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
        elif ext in [".xlsx", ".xls"]:
            wb = openpyxl.load_workbook(file_path)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    text += " | ".join([str(cell) if cell is not None else "" for cell in row]) + "\n"
        elif ext == ".docx":
            doc = Document(file_path)
            for para in doc.paragraphs:
                text += para.text + "\n"
    except Exception as e:
        print("Erreur extraction:", e)
    return text.strip()

# ----------------- Routes -----------------
@app.post("/upload/")
async def upload(file: UploadFile = File(...)):
    try:
        file_path = save_file(file)
        extracted_text = extract_text(file_path)
        lower_name = file.filename.lower()
        if "facture" in lower_name:
            collection = factures_col
            doc_type = "facture"
        elif "bdc" in lower_name:
            collection = bdcs_col
            doc_type = "bdc"
        else:
            return JSONResponse(status_code=400, content={"error": "Nom de fichier invalide (Facture/BDC attendu)"})

        doc = {
            "filename": file.filename,
            "filepath": file_path,
            "mimetype": file.content_type,
            "created_at": datetime.utcnow(),
            "extracted_text": extracted_text
        }
        res = collection.insert_one(doc)
        doc["_id"] = str(res.inserted_id)
        return {"message": "Upload réussi", "id": doc["_id"], "type": doc_type, "extracted_text": extracted_text}
    except Exception as e:
        return JSONResponse(status_code=500, content={"error": str(e)})

@app.get("/documents")
async def list_documents():
    factures = list(factures_col.find())
    bdcs = list(bdcs_col.find())
    for d in factures + bdcs:
        d["_id"] = str(d["_id"])
    return {"factures": factures, "bdcs": bdcs}

@app.get("/download/{doc_type}/{doc_id}")
async def download(doc_type: str, doc_id: str):
    collection = factures_col if doc_type == "facture" else bdcs_col
    doc = collection.find_one({"_id": ObjectId(doc_id)})
    if not doc:
        return JSONResponse(status_code=404, content={"error": "Document introuvable"})
    if not os.path.exists(doc["filepath"]):
        return JSONResponse(status_code=404, content={"error": "Fichier introuvable sur serveur"})
    return FileResponse(path=doc["filepath"], filename=doc["filename"], media_type=doc["mimetype"])

@app.post("/update/{doc_id}")
async def update_document(doc_id: str, modified_text: dict):
    text = modified_text.get("modified_text", "")
    for collection in [factures_col, bdcs_col]:
        res = collection.update_one({"_id": ObjectId(doc_id)}, {"$set": {"extracted_text": text, "updated_at": datetime.utcnow()}})
        if res.modified_count == 1:
            return {"message": "Document mis à jour avec succès"}
    return {"message": "Document non trouvé ou aucune modification"}

@app.delete("/delete/{doc_type}/{doc_id}")
async def delete_document(doc_type: str, doc_id: str):
    collection = factures_col if doc_type == "facture" else bdcs_col
    doc = collection.find_one({"_id": ObjectId(doc_id)})
    if not doc:
        return JSONResponse(status_code=404, content={"error": "Document introuvable"})
    if os.path.exists(doc["filepath"]):
        os.remove(doc["filepath"])
    collection.delete_one({"_id": ObjectId(doc_id)})
    return {"message": f"{doc_type.capitalize()} supprimé avec succès"}
